#!/usr/bin/env python3
"""Generate a polished Ctrl+Teach / Tars hackathon PowerPoint deck.

Dependencies are intentionally loaded from a temp target so this script does
not require modifying the app dependency graph:

    python3 -m pip install --target /private/tmp/ctrlteach-ppt-deps python-pptx pillow lxml
    python3 scripts/create_tars_pitch_deck_pro.py
"""

from __future__ import annotations

import math
import os
import sys
from pathlib import Path

sys.path.insert(0, "/private/tmp/ctrlteach-ppt-deps")

from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("CtrlTeach_Tars_Hackathon_Pitch_FINAL.pptx")
ASSET_DIR = Path("/private/tmp/ctrlteach-pitch-assets")
ASSET_DIR.mkdir(parents=True, exist_ok=True)


WIDE_W = Inches(13.333333)
WIDE_H = Inches(7.5)


class C:
    NAVY = "08111F"
    NAVY2 = "0F172A"
    INK = "172033"
    SLATE = "334155"
    MUTED = "64748B"
    SOFT = "F8FAFC"
    CREAM = "FFF7ED"
    WHITE = "FFFFFF"
    BLUE = "2563EB"
    BLUE2 = "38BDF8"
    CYAN = "22D3EE"
    GREEN = "22C55E"
    ORANGE = "F97316"
    AMBER = "F59E0B"
    PURPLE = "8B5CF6"
    ROSE = "FB7185"
    LINE = "CBD5E1"
    LAV = "EEF2FF"
    SKY = "E0F2FE"
    PEACH = "FFEDD5"


def rgb(hex_color: str) -> RGBColor:
    hex_color = hex_color.strip("#")
    return RGBColor(int(hex_color[0:2], 16), int(hex_color[2:4], 16), int(hex_color[4:6], 16))


def hex_tuple(hex_color: str) -> tuple[int, int, int]:
    hex_color = hex_color.strip("#")
    return (int(hex_color[:2], 16), int(hex_color[2:4], 16), int(hex_color[4:], 16))


def font_path() -> str | None:
    candidates = [
        "/System/Library/Fonts/Supplemental/Avenir Next.ttc",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
    ]
    return next((p for p in candidates if Path(p).exists()), None)


FONT_PATH = font_path()


def pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    if FONT_PATH:
        for index in ([1, 0] if bold else [0, 1]):
            try:
                return ImageFont.truetype(FONT_PATH, size=size, index=index)
            except OSError:
                continue
    return ImageFont.load_default()


def gradient_png(
    name: str,
    size: tuple[int, int],
    top: str,
    bottom: str,
    glow: list[tuple[float, float, int, str, int]] | None = None,
) -> Path:
    w, h = size
    top_rgb, bot_rgb = hex_tuple(top), hex_tuple(bottom)
    img = Image.new("RGB", size)
    px = img.load()
    for y in range(h):
        t = y / max(1, h - 1)
        color = tuple(int(top_rgb[i] * (1 - t) + bot_rgb[i] * t) for i in range(3))
        for x in range(w):
            px[x, y] = color
    img = img.convert("RGBA")
    if glow:
        layer = Image.new("RGBA", size, (0, 0, 0, 0))
        draw = ImageDraw.Draw(layer)
        for cx, cy, radius, color, alpha in glow:
            x = int(cx * w)
            y = int(cy * h)
            r = radius
            draw.ellipse((x - r, y - r, x + r, y + r), fill=(*hex_tuple(color), alpha))
        layer = layer.filter(ImageFilter.GaussianBlur(65))
        img = Image.alpha_composite(img, layer)
    out = ASSET_DIR / name
    img.save(out)
    return out


def browser_mock_png(name: str, label: str, callouts: list[str], mode: str = "aws") -> Path:
    w, h = 1400, 900
    img = Image.new("RGBA", (w, h), (248, 250, 252, 255))
    d = ImageDraw.Draw(img)

    # Browser frame
    d.rounded_rectangle((28, 28, w - 28, h - 28), radius=34, fill=(255, 255, 255), outline=(203, 213, 225), width=3)
    d.rounded_rectangle((28, 28, w - 28, 112), radius=34, fill=(241, 245, 249), outline=(203, 213, 225), width=3)
    d.rectangle((28, 78, w - 28, 112), fill=(241, 245, 249))
    for i, col in enumerate([(251, 113, 133), (245, 158, 11), (34, 197, 94)]):
        d.ellipse((65 + i * 34, 58, 86 + i * 34, 79), fill=col)
    d.rounded_rectangle((190, 52, 950, 88), radius=16, fill=(255, 255, 255), outline=(226, 232, 240), width=2)
    d.text((215, 60), label, font=pil_font(22, True), fill=hex_tuple(C.SLATE))

    # Left nav
    d.rectangle((28, 112, 268, h - 28), fill=(15, 23, 42))
    d.text((62, 152), "AWS Console", font=pil_font(32, True), fill=(255, 255, 255))
    nav = ["EC2", "Instances", "Security Groups", "Key Pairs", "Load Balancers", "Elastic IPs"]
    for i, item in enumerate(nav):
        y = 220 + i * 62
        fill = (37, 99, 235) if i == 1 else (30, 41, 59)
        d.rounded_rectangle((52, y, 240, y + 42), radius=12, fill=fill)
        d.text((72, y + 10), item, font=pil_font(20), fill=(255, 255, 255))

    # Main content
    d.text((315, 155), "Instances", font=pil_font(44, True), fill=hex_tuple(C.INK))
    d.text((318, 210), "Launch and manage EC2 instances", font=pil_font(24), fill=hex_tuple(C.MUTED))
    d.rounded_rectangle((1110, 150, 1300, 206), radius=16, fill=hex_tuple(C.ORANGE))
    d.text((1132, 166), "Launch instance", font=pil_font(23, True), fill=(255, 255, 255))
    d.rounded_rectangle((315, 260, 1300, 760), radius=24, fill=(255, 255, 255), outline=(226, 232, 240), width=3)
    for y in [335, 410, 485, 560, 635]:
        d.line((350, y, 1262, y), fill=(226, 232, 240), width=2)
    for x, t in [(355, "Name"), (620, "Instance ID"), (860, "State"), (1035, "Type")]:
        d.text((x, 290), t, font=pil_font(21, True), fill=hex_tuple(C.SLATE))
    rows = [
        ("prod-api", "i-0a42c9...", "running", "t3.micro"),
        ("demo-web", "i-0319ab...", "stopped", "t2.small"),
        ("training-box", "i-0ff47d...", "running", "t3.medium"),
    ]
    for i, row in enumerate(rows):
        y = 362 + i * 74
        for x, t in zip([355, 620, 860, 1035], row):
            d.text((x, y), t, font=pil_font(23), fill=hex_tuple(C.SLATE))

    # Tars bubble and pointer
    d.rounded_rectangle((778, 52, 1340, 128), radius=32, fill=(8, 17, 31, 235))
    d.ellipse((805, 72, 845, 112), fill=hex_tuple(C.CYAN))
    d.text((860, 74), "Hold Ctrl and ask: “where do I click?”", font=pil_font(24, True), fill=(255, 255, 255))
    d.line((1090, 130, 1180, 154), fill=hex_tuple(C.CYAN), width=8)
    d.ellipse((1158, 132, 1210, 184), outline=hex_tuple(C.CYAN), width=7)

    # Optional capability callouts
    colors = [C.BLUE, C.PURPLE, C.GREEN]
    for i, text in enumerate(callouts[:3]):
        x = 320 + i * 318
        y = 792
        d.rounded_rectangle((x, y, x + 285, y + 58), radius=20, fill=hex_tuple(colors[i]))
        d.text((x + 22, y + 16), text, font=pil_font(22, True), fill=(255, 255, 255))

    out = ASSET_DIR / name
    img.save(out)
    return out


def classroom_mock_png(name: str) -> Path:
    w, h = 1400, 900
    img = Image.new("RGBA", (w, h), hex_tuple(C.SOFT) + (255,))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((42, 42, w - 42, h - 42), radius=36, fill=(255, 255, 255), outline=(203, 213, 225), width=3)
    d.text((90, 85), "Live Classroom Mode", font=pil_font(46, True), fill=hex_tuple(C.INK))
    d.text((92, 142), "voice + whiteboard + screen annotations", font=pil_font(26), fill=hex_tuple(C.MUTED))
    d.rounded_rectangle((90, 210, 900, 765), radius=24, fill=(255, 247, 237), outline=(253, 186, 116), width=3)
    d.text((140, 255), "EC2 launch flow", font=pil_font(38, True), fill=hex_tuple(C.INK))
    points = [(190, 385), (390, 385), (590, 385), (790, 385)]
    labels = ["AMI", "Type", "Key pair", "Launch"]
    for i, (x, y) in enumerate(points):
        d.ellipse((x - 48, y - 48, x + 48, y + 48), fill=hex_tuple([C.BLUE, C.PURPLE, C.ORANGE, C.GREEN][i]))
        d.text((x - 18, y - 18), str(i + 1), font=pil_font(34, True), fill=(255, 255, 255))
        d.text((x - 48, y + 64), labels[i], font=pil_font(24, True), fill=hex_tuple(C.SLATE))
        if i < len(points) - 1:
            d.line((x + 54, y, points[i + 1][0] - 54, y), fill=hex_tuple(C.SLATE), width=5)
    # Drawing squiggle
    pts = [(165, 560), (250, 520), (345, 575), (440, 530), (540, 590), (670, 540), (790, 610)]
    d.line(pts, fill=hex_tuple(C.ORANGE), width=8, joint="curve")
    d.text((145, 640), "Tars can draw and explain live", font=pil_font(30, True), fill=hex_tuple(C.ORANGE))
    # Chat panel
    d.rounded_rectangle((960, 210, 1305, 765), radius=24, fill=hex_tuple(C.NAVY))
    d.ellipse((1005, 260, 1065, 320), fill=hex_tuple(C.CYAN))
    d.text((1085, 270), "Tars", font=pil_font(30, True), fill=(255, 255, 255))
    bubbles = [
        ("Tell me what you’re trying to deploy.", 270),
        ("I’ll draw the path while we talk.", 395),
        ("Now click Launch instance.", 520),
    ]
    for text, y in bubbles:
        d.rounded_rectangle((1005, y + 70, 1265, y + 136), radius=20, fill=(30, 41, 59))
        d.text((1025, y + 88), text, font=pil_font(22), fill=(226, 232, 240))
    out = ASSET_DIR / name
    img.save(out)
    return out


def add_bg(slide, path: Path):
    slide.shapes.add_picture(str(path), 0, 0, width=WIDE_W, height=WIDE_H)


def add_rect(slide, x, y, w, h, fill, line=None, radius=True, transparency=0):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = rgb(fill)
    shp.fill.transparency = transparency
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    return shp


def add_text(
    slide,
    text,
    x,
    y,
    w,
    h,
    size=24,
    color=C.INK,
    bold=False,
    align=PP_ALIGN.LEFT,
    font="Avenir Next",
    margin=0.08,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_multiline(slide, lines, x, y, w, h, size=21, color=C.SLATE, gap=0.34):
    y0 = y
    for i, line in enumerate(lines):
        add_text(slide, line, x, y0 + i * gap, w, 0.32, size=size, color=color)


def add_title(slide, title, subtitle=None, dark=False):
    add_text(slide, title, 0.62, 0.38, 11.8, 0.55, size=30, color=C.WHITE if dark else C.INK, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.65, 0.98, 11.4, 0.55, size=15.5, color="CBD5E1" if dark else C.MUTED)


def add_pill(slide, text, x, y, w, fill=C.BLUE, color=C.WHITE):
    add_rect(slide, x, y, w, 0.36, fill, radius=True)
    add_text(slide, text, x + 0.04, y + 0.04, w - 0.08, 0.22, size=10.5, color=color, bold=True, align=PP_ALIGN.CENTER, margin=0)


def add_placeholder(slide, title, sub, x, y, w, h):
    add_rect(slide, x, y, w, h, C.SOFT, C.LINE, radius=True)
    title_size = 13 if w < 3.2 else 15
    title_y = y + h / 2 - 0.55
    add_text(
        slide,
        title,
        x + 0.25,
        title_y,
        w - 0.5,
        0.72,
        size=title_size,
        color=C.BLUE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_text(
        slide,
        sub,
        x + 0.38,
        title_y + 0.78,
        w - 0.76,
        0.42,
        size=10,
        color=C.MUTED,
        align=PP_ALIGN.CENTER,
    )


def add_card(slide, title, body, x, y, w, h, accent=C.BLUE):
    add_rect(slide, x, y, w, h, C.WHITE, "E2E8F0", radius=True)
    add_rect(slide, x, y, 0.08, h, accent, radius=False)
    add_text(slide, title, x + 0.24, y + 0.22, w - 0.45, 0.38, size=17, color=accent, bold=True)
    add_text(slide, body, x + 0.24, y + 0.74, w - 0.42, h - 0.9, size=12.5, color=C.SLATE)


def add_bullets(slide, bullets, x, y, w, size=14.5, color=C.SLATE, bullet_color=C.BLUE, line_gap=0.48):
    for i, b in enumerate(bullets):
        yy = y + i * line_gap
        add_rect(slide, x, yy + 0.06, 0.13, 0.13, bullet_color, radius=True)
        add_text(slide, b, x + 0.26, yy - 0.02, w - 0.26, 0.34, size=size, color=color)


def image_with_shadow(slide, path, x, y, w, h):
    add_rect(slide, x + 0.08, y + 0.08, w, h, "000000", radius=True, transparency=78)
    slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def build_deck():
    prs = Presentation()
    prs.slide_width = WIDE_W
    prs.slide_height = WIDE_H
    blank = prs.slide_layouts[6]

    bg_dark = gradient_png(
        "bg_dark.png",
        (1920, 1080),
        C.NAVY,
        "111827",
        glow=[(0.88, 0.12, 430, C.CYAN, 90), (0.08, 0.82, 420, C.PURPLE, 90), (0.72, 0.92, 360, C.BLUE, 70)],
    )
    bg_light = gradient_png(
        "bg_light.png",
        (1920, 1080),
        "FFF7ED",
        "EFF6FF",
        glow=[(0.85, 0.1, 330, C.CYAN, 35), (0.13, 0.85, 360, C.ORANGE, 35)],
    )
    bg_blue = gradient_png(
        "bg_blue.png",
        (1920, 1080),
        "E0F2FE",
        "EEF2FF",
        glow=[(0.72, 0.18, 340, C.BLUE, 45), (0.18, 0.82, 330, C.PURPLE, 38)],
    )
    aws_img = browser_mock_png(
        "aws_tars_mock.png",
        "https://console.aws.amazon.com/ec2/",
        ["points", "draws", "guides"],
    )
    lab_img = browser_mock_png(
        "lab_mock.png",
        "Ctrl+Teach browser lab — AWS EC2 deployment",
        ["tracks", "evidence", "cleanup"],
    )
    classroom_img = classroom_mock_png("classroom_mock.png")

    # 1 Cover
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_dark)
    add_pill(slide, "Cprime's next scalable product", 0.78, 0.62, 2.55, C.BLUE2)
    add_text(slide, "Ctrl+Teach", 0.72, 1.28, 6.2, 0.72, size=47, color=C.WHITE, bold=True)
    add_text(slide, "Cprime's expertise,", 0.74, 2.06, 6.0, 0.62, size=34, color=C.WHITE, bold=True)
    add_text(slide, "turned into a product.", 0.74, 2.64, 6.0, 0.62, size=34, color=C.CYAN, bold=True)
    add_text(
        slide,
        "Personalized courses on demand — with a tutor that stays beside every learner until they can do the work.",
        0.78,
        3.48,
        5.95,
        0.9,
        size=17.5,
        color="DDEBFF",
    )
    add_rect(slide, 0.78, 4.72, 5.65, 1.18, "172033", "334155", radius=True, transparency=5)
    add_text(
        slide,
        "Build once. Improve continuously. Sell repeatedly.\nBut that’s not the best part — every learner gets Tars.",
        1.0,
        4.98,
        5.2,
        0.55,
        size=15,
        color=C.CREAM,
        bold=True,
    )
    add_placeholder(
        slide,
        "HERO SCREENSHOT: TARS BESIDE THE CURSOR",
        "Paste your strongest product screenshot here",
        7.0,
        1.0,
        5.6,
        5.15,
    )

    # 2 Tension / thesis
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_light)
    add_title(
        slide,
        "The uncomfortable bottleneck",
        "Cprime's expertise is valuable. But when every new cohort needs another trainer, schedule, and delivery cycle, growth stays capped.",
    )
    add_text(slide, "Today: expertise tied to hours", 0.9, 2.05, 4.65, 0.42, size=22, color=C.ROSE, bold=True)
    add_text(slide, "With Ctrl+Teach: expertise becomes product", 7.35, 2.05, 5.15, 0.42, size=22, color=C.BLUE, bold=True)
    add_rect(slide, 0.82, 2.75, 4.55, 2.65, C.WHITE, "E2E8F0", radius=True)
    add_bullets(
        slide,
        ["Growth depends on trainer availability", "The same expertise is delivered again", "Tools change before courses catch up", "Group classes split individual attention"],
        1.15,
        3.15,
        3.8,
        bullet_color=C.ROSE,
    )
    add_text(slide, "→", 6.13, 3.48, 0.8, 0.6, size=45, color=C.BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 7.1, 2.75, 4.95, 2.65, C.WHITE, "BFDBFE", radius=True)
    add_bullets(
        slide,
        ["One learning product reaches thousands", "Update once and everyone benefits", "Every learner gets personal support", "Value grows without matching headcount"],
        7.45,
        3.15,
        4.15,
        bullet_color=C.BLUE,
    )
    add_rect(slide, 1.55, 6.15, 10.25, 0.65, C.NAVY, radius=True)
    add_text(
        slide,
        "Every new learner should create leverage — not another delivery burden.",
        1.7,
        6.32,
        9.95,
        0.28,
        size=19,
        color=C.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    # 3 Course engine / leverage
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_blue)
    add_title(
        slide,
        "Build once. Improve continuously. Sell repeatedly.",
        "Ctrl+Teach turns Cprime's learning expertise into software that becomes more useful with every learner.",
    )
    add_card(slide, "Generated on demand", "Create complete courses around the learner's actual goal — without weeks of manual production.", 0.72, 1.72, 2.75, 1.52, C.BLUE)
    add_card(slide, "Beautiful and credible", "Rich visuals and trustworthy sources make the content feel polished, current, and worth learning from.", 3.66, 1.72, 2.75, 1.52, C.ORANGE)
    add_card(slide, "Built for practice", "Quizzes, examples, and real browser labs move learners from knowing to doing.", 0.72, 3.48, 2.75, 1.52, C.PURPLE)
    add_card(slide, "Personal over time", "Learner memories shape the next explanation, analogy, and course around how that person thinks.", 3.66, 3.48, 2.75, 1.52, C.GREEN)
    add_placeholder(
        slide,
        "SCREENSHOT: GENERATED COURSE",
        "Paste a course page showing visuals, sources, quizzes, or labs",
        6.85,
        1.7,
        5.75,
        4.75,
    )
    add_rect(slide, 0.92, 5.55, 5.15, 0.9, C.NAVY, radius=True)
    add_text(
        slide,
        "One engine. Thousands of personalized learning journeys.",
        1.12,
        5.82,
        4.75,
        0.28,
        size=15.5,
        color=C.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    # 4 AWS story
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_blue)
    add_title(
        slide,
        "The AWS moment we all remember",
        "Day one: create an EC2 instance and deploy the app. The trainer showed it once — then we had to do it ourselves.",
    )
    add_card(slide, "What happened", "We forgot where to click and kept jumping between ChatGPT and AWS.", 0.78, 1.82, 4.1, 1.3, C.ORANGE)
    add_card(slide, "Why normal AI failed", "It said the button was “top left.” AWS had changed the UI. We were still stuck.", 0.78, 3.38, 4.1, 1.65, C.ROSE)
    add_card(slide, "The Tars version", "Hold Ctrl, point at the screen, and ask. Tars sees this exact screen and points to the right place.", 0.78, 5.28, 4.1, 1.32, C.BLUE)
    add_placeholder(
        slide,
        "SCREENSHOT: TARS POINTING INSIDE AWS",
        "Paste the real EC2 guidance moment here",
        5.45,
        1.66,
        6.95,
        4.98,
    )

    # 5 Meet Tars
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_dark)
    add_title(
        slide,
        "But that’s not the best part.",
        "Every course comes with Tars — a small AI tutor that lives beside your cursor.",
        dark=True,
    )
    add_rect(slide, 0.85, 1.85, 4.8, 4.7, "111827", "334155", radius=True, transparency=0)
    add_text(slide, "When the learner gets stuck:", 1.12, 2.18, 4.1, 0.35, size=20, color=C.WHITE, bold=True)
    add_bullets(
        slide,
        ["hold Ctrl", "point, circle, or underline", "ask out loud or type", "Tars answers in context"],
        1.18,
        2.82,
        3.8,
        color="DDEBFF",
        bullet_color=C.CYAN,
        size=15.2,
    )
    add_rect(slide, 1.18, 5.2, 3.9, 0.72, C.BLUE, radius=True)
    add_text(slide, "“Where do I click now?”", 1.28, 5.4, 3.7, 0.22, size=16, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_placeholder(slide, "SCREENSHOT: TARS BESIDE THE CURSOR", "Paste your floating Tars UI here", 6.15, 1.75, 5.95, 4.95)

    # 6 Capabilities
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_light)
    add_title(slide, "Tars doesn’t just answer. It sees, shows, and acts.", "Use the four spaces below for real product screenshots — one clear capability per frame.")
    caps = [
        ("Points exactly where to click", "No vague directions. Tars grounds the answer to the learner’s exact screen.", "POINTING SCREENSHOT", C.BLUE),
        ("Draws over the screen", "It sketches, underlines, and circles directly over the browser.", "DRAWING SCREENSHOT", C.ORANGE),
        ("Understands what you circle", "Hold Ctrl and mark a region. Tars knows what you mean.", "CTRL + CIRCLE SCREENSHOT", C.PURPLE),
        ("Takes browser actions", "It can scroll, switch tabs, navigate, and guide multi-step workflows.", "BROWSER ACTION SCREENSHOT", C.GREEN),
    ]
    for i, (title, body, placeholder_title, col) in enumerate(caps):
        x = 0.68 + (i % 2) * 6.35
        y = 1.62 + (i // 2) * 2.63
        add_rect(slide, x, y, 5.95, 2.35, C.WHITE, "E2E8F0", radius=True)
        add_rect(slide, x, y, 0.08, 2.35, col, radius=False)
        add_text(slide, title, x + 0.25, y + 0.25, 2.38, 0.52, size=16.5, color=col, bold=True)
        add_text(slide, body, x + 0.25, y + 0.92, 2.38, 0.95, size=11.5, color=C.SLATE)
        add_placeholder(slide, placeholder_title, "Paste image here", x + 2.88, y + 0.25, 2.75, 1.84)

    # 7 Ctrl help detail
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_dark)
    add_title(slide, "The Ctrl gesture is the magic interaction", "Instead of describing the UI in words, the learner can simply refer to the screen.", dark=True)
    add_rect(slide, 0.82, 1.75, 5.25, 4.95, "111827", "334155", radius=True)
    add_text(slide, "Hold Ctrl + point", 1.15, 2.05, 4.4, 0.4, size=24, color=C.CYAN, bold=True)
    add_text(
        slide,
        "Tars understands the exact part of the screen you mean and answers in context — no screenshots, no guessing, no stale directions.",
        1.15,
        2.68,
        4.45,
        1.05,
        size=15.5,
        color="E2E8F0",
    )
    add_text(slide, "Why it matters:", 1.15, 4.02, 4.2, 0.3, size=17, color=C.WHITE, bold=True)
    add_bullets(
        slide,
        ["No screenshot upload loop", "No stale “top left button” advice", "No guessing what the learner meant"],
        1.2,
        4.55,
        4.2,
        color="DDEBFF",
        bullet_color=C.CYAN,
        size=14,
    )
    add_placeholder(slide, "SCREENSHOT: CTRL TRAIL + REGION", "Show the fading trail and the selected area", 6.75, 1.75, 5.45, 4.95)

    # 8 Browser labs
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_light)
    add_title(slide, "Learning does not stop when the lesson ends", "Tars stays while the learner implements what they just learned inside the real platform.")
    add_placeholder(slide, "SCREENSHOT: REAL BROWSER LAB", "Paste the AWS, Jira, GitHub, ServiceNow, or Power BI lab here", 0.78, 1.65, 6.6, 4.75)
    add_card(slide, "Practice for real", "Open the actual platform and complete the real workflow — not a fake simulation.", 7.8, 1.78, 4.55, 1.35, C.ORANGE)
    add_card(slide, "Tars watches closely", "It monitors progress, steps in when needed, and never makes mistakes feel embarrassing.", 7.8, 3.42, 4.55, 1.35, C.BLUE)
    add_card(slide, "Completion means proven", "The backend confirms the task, evidence, and required cleanup before marking it complete.", 7.8, 5.06, 4.55, 1.35, C.GREEN)

    # 9 Classroom mode
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_blue)
    add_title(slide, "A live classroom for one learner at a time", "It feels like a real tutor — but every learner gets the attention a group class cannot provide.")
    add_rect(slide, 0.68, 1.65, 3.38, 5.15, C.WHITE, "BFDBFE", radius=True)
    add_text(slide, "Two ways to learn", 1.02, 1.98, 2.7, 0.34, size=20, color=C.INK, bold=True, align=PP_ALIGN.CENTER)
    add_card(slide, "Regular text mode", "Read, review, take quizzes, and ask quick questions.", 0.98, 2.63, 2.78, 1.28, C.BLUE)
    add_card(slide, "Live Classroom Mode", "Talk naturally while Tars explains and draws diagrams on a whiteboard.", 0.98, 4.12, 2.78, 1.48, C.ORANGE)
    add_text(slide, "Then it watches you apply the lesson, helps when you slip, and never judges the mistake.", 1.02, 5.92, 2.72, 0.62, size=11.5, color=C.SLATE, bold=True, align=PP_ALIGN.CENTER)
    add_placeholder(slide, "SCREENSHOT: LIVE CLASSROOM + WHITEBOARD", "Paste Tars teaching or drawing on the whiteboard here", 4.42, 1.65, 7.9, 5.15)

    # 10 Proof / architecture
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_light)
    add_title(slide, "Not “they attended.” Proof they can do the work.", "Every lab produces evidence that the learner completed the task correctly — including required cleanup.")
    flow = [
        ("Learn", "Understand the concept", C.BLUE),
        ("Practice", "Work in the real tool", C.ORANGE),
        ("Tars", "Observes implementation", C.PURPLE),
        ("Verify", "Confirm task + cleanup", C.GREEN),
    ]
    for i, (head, body, col) in enumerate(flow):
        x = 0.72 + i * 3.14
        add_rect(slide, x, 2.12, 2.45, 2.25, C.WHITE, "E2E8F0", radius=True)
        add_rect(slide, x + 0.75, 2.42, 0.95, 0.95, col, radius=True)
        add_text(slide, str(i + 1), x + 0.86, 2.58, 0.75, 0.32, size=24, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, head, x + 0.2, 3.58, 2.05, 0.25, size=17, color=col, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.22, 3.93, 2.0, 0.35, size=10.5, color=C.SLATE, align=PP_ALIGN.CENTER)
        if i < 3:
            add_text(slide, "→", x + 2.52, 2.85, 0.45, 0.4, size=30, color=C.MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_rect(slide, 1.1, 5.52, 11.15, 0.72, C.NAVY, radius=True)
    add_text(slide, "The learner does not just finish the course. They prove they can perform the workflow.", 1.25, 5.72, 10.85, 0.28, size=17.5, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)

    # 11 Impact
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_dark)
    add_title(slide, "This is Cprime’s product leverage", "Build the learning experience once. Improve it continuously. Sell it again and again.", dark=True)
    impact = [
        ("Revenue without matching headcount", "Serve the next 1,000 learners without finding 1,000 more trainer-hours.", C.GREEN),
        ("Expertise becomes an asset", "Cprime’s knowledge keeps earning instead of restarting with every cohort.", C.CYAN),
        ("Personalization at scale", "Every learner gets individual attention, memory, and relevant examples.", C.ORANGE),
        ("A true scalable product", "This could become Cprime’s first learning product sold repeatedly like software.", C.PURPLE),
    ]
    for i, (t, b, col) in enumerate(impact):
        add_rect(slide, 0.82 + (i % 2) * 6.02, 1.85 + (i // 2) * 2.0, 5.45, 1.42, "111827", "334155", radius=True)
        add_text(slide, t, 1.12 + (i % 2) * 6.02, 2.14 + (i // 2) * 2.0, 4.9, 0.3, size=19, color=col, bold=True)
        add_text(slide, b, 1.12 + (i % 2) * 6.02, 2.62 + (i // 2) * 2.0, 4.85, 0.48, size=13.5, color="DDEBFF")
    add_text(slide, "Cprime stops scaling only through trainer time — and starts scaling through a product it owns.", 1.15, 6.15, 11.1, 0.55, size=20, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)

    # 12 Demo story
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_light)
    add_title(slide, "Demo the transformation in 90 seconds", "Show one learner moving from stuck to capable — without ever leaving the real tool.")
    demo_frames = [
        ("1. LEARNER GETS STUCK", "Show the confusing AWS screen", C.ORANGE),
        ("2. CTRL + TARS", "Show Tars pointing or drawing", C.PURPLE),
        ("3. TASK COMPLETED", "Show success or backend verification", C.GREEN),
    ]
    for i, (title, sub, col) in enumerate(demo_frames):
        x = 0.64 + i * 4.25
        add_placeholder(slide, title, sub, x, 1.72, 3.85, 3.8)
        add_rect(slide, x + 0.65, 5.76, 2.55, 0.42, col, radius=True)
        add_text(slide, ["The confusion", "The wow moment", "The proof"][i], x + 0.73, 5.86, 2.38, 0.2, size=11.5, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_pill(slide, "One story. One clear before-and-after.", 4.32, 6.62, 4.7, C.BLUE)

    # 13 Close
    slide = prs.slides.add_slide(blank)
    add_bg(slide, bg_dark)
    add_text(slide, "Ctrl+Teach", 0.75, 0.62, 4.0, 0.45, size=26, color=C.CYAN, bold=True)
    add_text(slide, "The future of learning is not", 0.95, 1.5, 11.45, 0.58, size=32, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "another video library.", 0.95, 2.05, 11.45, 0.58, size=32, color=C.ROSE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "It is a tutor that stays until you can do it yourself.", 0.9, 2.75, 11.55, 0.62, size=28, color=C.CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Built once. Personalized endlessly. Delivered at software scale.", 1.35, 3.48, 10.65, 0.42, size=18, color="CBD5E1", align=PP_ALIGN.CENTER)
    add_placeholder(slide, "FINAL PRODUCT SCREENSHOT", "Use your strongest Tars visual here", 3.4, 4.28, 6.55, 1.72)
    add_pill(slide, "Ctrl+Teach + Tars", 5.15, 6.55, 3.0, C.BLUE2)

    # Metadata-ish properties
    prs.core_properties.title = "Ctrl+Teach — Cprime's Expertise, Turned Into a Product"
    prs.core_properties.subject = "Hackathon pitch deck"
    prs.core_properties.author = "Ctrl+Teach"
    prs.save(OUT)
    print(f"Wrote {OUT.resolve()} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build_deck()
