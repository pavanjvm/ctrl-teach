#!/usr/bin/env python3
"""Create a clean, content-first Ctrl+Teach hackathon pitch deck."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, "/private/tmp/ctrlteach-ppt-deps")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path("CtrlTeach_Hackathon_Pitch_CLEAN_MODERN.pptx")

SLIDE_W = Inches(13.333333)
SLIDE_H = Inches(7.5)


class C:
    INK = "101828"
    NAVY = "0B1220"
    WHITE = "FFFFFF"
    BG = "F7F8FA"
    SOFT = "F2F4F7"
    SLATE = "475467"
    MUTED = "667085"
    LINE = "D0D5DD"
    BLUE = "2563EB"
    BLUE_DARK = "1D4ED8"
    BLUE_SOFT = "EFF6FF"
    TEAL = "0D9488"
    TEAL_SOFT = "ECFDF5"
    AMBER = "F59E0B"
    AMBER_SOFT = "FFFAEB"
    DARK_LINE = "344054"
    DARK_SOFT = "182230"
    LIGHT_TEXT = "E4E7EC"


FONT = "Avenir Next"


def rgb(value: str) -> RGBColor:
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def set_background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: str | None = None,
    line: str | None = None,
    radius: bool = False,
    line_width: float = 1,
    transparency: int = 0,
):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
        shape.fill.transparency = transparency
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(line_width)
    else:
        shape.line.fill.background()
    return shape


def add_circle(slide, x: float, y: float, d: float, *, fill: str, line: str | None = None):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = C.INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0,
    font: str = FONT,
    italic: bool = False,
    spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    if spacing is not None:
        paragraph.line_spacing = spacing
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return box


def add_rule(slide, x: float, y: float, w: float, color: str = C.LINE, height: float = 0.012):
    add_box(slide, x, y, w, height, fill=color)


def add_header(slide, section: str, title: str, subtitle: str | None, page: int, *, dark: bool = False):
    title_color = C.WHITE if dark else C.INK
    subtitle_color = C.LIGHT_TEXT if dark else C.SLATE
    add_box(slide, 0.78, 0.43, 0.055, 0.24, fill=C.BLUE)
    add_text(slide, section.upper(), 0.96, 0.41, 7.5, 0.25, size=9.5, color=C.BLUE if not dark else "93C5FD", bold=True)
    add_text(slide, title, 0.78, 0.79, 11.75, 0.56, size=30, color=title_color, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.78, 1.38, 11.7, 0.5, size=14.5, color=subtitle_color)
    add_text(slide, f"{page:02d}", 12.0, 0.43, 0.55, 0.22, size=9.5, color=C.MUTED if not dark else "98A2B3", bold=True, align=PP_ALIGN.RIGHT)


def add_footer(slide, page: int, *, dark: bool = False):
    line = C.DARK_LINE if dark else C.LINE
    text_color = "98A2B3" if dark else C.MUTED
    add_rule(slide, 0.78, 7.05, 11.77, line)
    add_text(slide, "CTRL+TEACH  /  HACKATHON CONCEPT", 0.78, 7.14, 4.5, 0.18, size=8.5, color=text_color, bold=True)
    add_text(slide, f"{page:02d}", 12.0, 7.14, 0.55, 0.18, size=8.5, color=text_color, bold=True, align=PP_ALIGN.RIGHT)


def add_browser_placeholder(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    label: str,
    caption: str,
    *,
    dark: bool = False,
):
    fill = C.DARK_SOFT if dark else C.WHITE
    top_fill = "202B3C" if dark else C.SOFT
    line = C.DARK_LINE if dark else C.LINE
    primary = C.LIGHT_TEXT if dark else C.SLATE
    secondary = "98A2B3" if dark else C.MUTED
    add_box(slide, x, y, w, h, fill=fill, line=line, radius=True, line_width=1)
    add_box(slide, x + 0.02, y + 0.02, w - 0.04, 0.34, fill=top_fill)
    for i, dot in enumerate(["98A2B3", "667085", C.BLUE]):
        add_circle(slide, x + 0.17 + i * 0.18, y + 0.125, 0.075, fill=dot)
    add_text(slide, "INSERT PRODUCT SCREENSHOT", x + 0.35, y + h / 2 - 0.34, w - 0.7, 0.22, size=9.5, color=secondary, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.38, y + h / 2 - 0.03, w - 0.76, 0.65, size=13, color=primary, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, caption, x, y + h + 0.12, w, 0.32, size=10, color=secondary)


def add_editorial_item(slide, number: str, title: str, body: str, x: float, y: float, w: float, *, accent: str = C.BLUE):
    add_text(slide, number, x, y, 0.42, 0.25, size=10.5, color=accent, bold=True)
    add_text(slide, title, x + 0.55, y - 0.02, w - 0.55, 0.3, size=15.5, color=C.INK, bold=True)
    add_text(slide, body, x + 0.55, y + 0.38, w - 0.55, 0.7, size=12.5, color=C.SLATE)


def add_dot_bullet(slide, text: str, x: float, y: float, w: float, *, color: str = C.BLUE, size: float = 13.5):
    add_circle(slide, x, y + 0.075, 0.09, fill=color)
    add_text(slide, text, x + 0.24, y, w - 0.24, 0.42, size=size, color=C.SLATE)


def build_deck() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # 1 — Cover
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.NAVY)
    add_box(slide, 0.78, 0.62, 0.06, 0.3, fill=C.BLUE)
    add_text(slide, "CTRL+TEACH  /  HACKATHON CONCEPT", 0.98, 0.62, 5.2, 0.28, size=10.5, color="93C5FD", bold=True)
    add_text(slide, "Cprime’s expertise.", 0.78, 1.55, 8.8, 0.7, size=42, color=C.WHITE, bold=True)
    add_text(slide, "Delivered at software scale.", 0.78, 2.28, 10.6, 0.72, size=42, color="60A5FA", bold=True)
    add_text(
        slide,
        "Personalized courses on demand — and a tutor that stays until the learner can do the work.",
        0.82,
        3.35,
        8.4,
        0.8,
        size=19,
        color=C.LIGHT_TEXT,
    )
    add_rule(slide, 0.82, 4.65, 8.75, C.DARK_LINE)
    cover_items = [
        ("BUILD ONCE", "Create the learning experience"),
        ("IMPROVE CONTINUOUSLY", "Keep content and guidance current"),
        ("SELL REPEATEDLY", "Reach more learners with the same expertise"),
    ]
    for i, (label, body) in enumerate(cover_items):
        x = 0.82 + i * 3.22
        if i:
            add_box(slide, x - 0.28, 5.02, 0.012, 0.92, fill=C.DARK_LINE)
        add_text(slide, label, x, 5.02, 2.8, 0.24, size=10, color="93C5FD", bold=True)
        add_text(slide, body, x, 5.42, 2.72, 0.58, size=12.5, color="B9C0CC")

    # 2 — Opportunity
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.WHITE)
    add_header(slide, "01  /  The opportunity", "Cprime already owns the hardest part.", None, 2)
    add_text(slide, "The expertise and trust to teach people.", 0.78, 1.62, 11.2, 0.5, size=25, color=C.BLUE, bold=True)
    add_text(
        slide,
        "Cprime has spent years building the expertise, instructors, partnerships, and course catalog that make great learning possible.",
        0.78,
        2.28,
        11.55,
        0.65,
        size=17,
        color=C.SLATE,
    )
    add_rule(slide, 0.78, 3.25, 11.77)
    columns = [
        ("01", "Already proven", "Deep domain expertise, trusted trainers, and a learning brand enterprises already know."),
        ("02", "A natural limit", "Reach still depends on trainer availability, group schedules, and delivery hours."),
        ("03", "The next layer", "Turn that knowledge into a product that reaches more people without losing the human feel."),
    ]
    for i, (num, title, body) in enumerate(columns):
        x = 0.78 + i * 4.0
        if i:
            add_box(slide, x - 0.28, 3.67, 0.012, 1.62, fill=C.LINE)
        add_text(slide, num, x, 3.62, 0.42, 0.24, size=10, color=C.BLUE, bold=True)
        add_text(slide, title, x, 4.02, 3.45, 0.34, size=17, color=C.INK, bold=True)
        add_text(slide, body, x, 4.55, 3.52, 0.88, size=13.5, color=C.SLATE)
    add_box(slide, 0.78, 5.82, 11.77, 0.82, fill=C.BLUE_SOFT)
    add_text(
        slide,
        "The opportunity is not to replace what works. It is to help Cprime’s knowledge travel further.",
        1.05,
        6.05,
        11.2,
        0.34,
        size=16,
        color=C.BLUE_DARK,
        bold=True,
    )
    add_footer(slide, 2)

    # 3 — Leverage
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.BG)
    add_header(slide, "02  /  The leverage", "Turn expertise into leverage.", "A learning product that gets better as it reaches more people.", 3)
    flow = [
        ("01", "Build once", "Create the course and tutor experience."),
        ("02", "Improve continuously", "Update once; every learner benefits."),
        ("03", "Personalize", "Memory automatically adapts the journey to the person."),
        ("04", "Sell repeatedly", "Reach new cohorts without recreating delivery."),
    ]
    for i, (num, title, body) in enumerate(flow):
        x = 0.78 + i * 3.0
        add_circle(slide, x, 2.08, 0.44, fill=C.BLUE)
        add_text(slide, num, x, 2.18, 0.44, 0.16, size=8.5, color=C.WHITE, bold=True, align=PP_ALIGN.CENTER)
        if i < 3:
            add_rule(slide, x + 0.58, 2.29, 2.28, C.LINE)
        add_text(slide, title, x, 2.82, 2.55, 0.34, size=16.5, color=C.INK, bold=True)
        add_text(slide, body, x, 3.35, 2.55, 0.72, size=12.5, color=C.SLATE)
    add_rule(slide, 0.78, 4.48, 11.77)
    add_text(
        slide,
        "Ctrl+Teach turns Cprime’s knowledge into a product it owns — one that can support thousands of learners without requiring thousands of additional delivery hours.",
        0.78,
        4.9,
        11.7,
        0.92,
        size=18,
        color=C.INK,
        bold=True,
    )
    add_text(slide, "One learning engine", 0.78, 6.15, 3.1, 0.28, size=12, color=C.MUTED, bold=True)
    add_text(slide, "→", 3.95, 6.12, 0.55, 0.3, size=18, color=C.BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Thousands of personalized journeys", 4.62, 6.15, 5.0, 0.28, size=12, color=C.BLUE, bold=True)
    add_footer(slide, 3)

    # 4 — Course generation and memory
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.WHITE)
    add_header(slide, "03  /  The course", "A complete course, generated on demand.", "Beautiful, credible, practical — and increasingly personal.", 4)
    add_text(
        slide,
        "Ctrl+Teach creates full learning journeys with beautiful visuals, credible sources, practical examples, quizzes, and hands-on browser labs.",
        0.78,
        1.93,
        7.15,
        0.82,
        size=17,
        color=C.INK,
        bold=True,
    )
    add_editorial_item(slide, "01", "Generated on demand", "Start from the learner’s real goal instead of a fixed catalog path.", 0.78, 3.02, 3.42)
    add_editorial_item(slide, "02", "Credible sources", "Give every lesson a trustworthy foundation and a clear trail back to it.", 4.24, 3.02, 3.48)
    add_editorial_item(slide, "03", "Built for practice", "Quizzes and real browser labs move the learner from knowing to doing.", 0.78, 4.45, 3.42)
    add_editorial_item(slide, "04", "Personal through memory", "The next course starts smarter because Ctrl+Teach remembers how the learner thinks.", 4.24, 4.45, 3.48, accent=C.TEAL)
    add_browser_placeholder(
        slide,
        8.35,
        1.93,
        4.18,
        3.0,
        "Personalized course or learner profile",
        "Show visuals, sources, or the learner-memory experience.",
    )
    add_box(slide, 8.35, 5.35, 4.18, 0.88, fill=C.TEAL_SOFT)
    add_text(slide, "A small example", 8.62, 5.57, 1.2, 0.2, size=9.5, color=C.TEAL, bold=True)
    add_text(slide, "If fitness examples click, the next lesson remembers and uses them again.", 9.83, 5.5, 2.38, 0.48, size=11.5, color=C.INK, bold=True)
    add_footer(slide, 4)

    # 5 — Tars reveal
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.NAVY)
    add_header(slide, "04  /  The tutor", "But that’s not the best part.", None, 5, dark=True)
    add_text(slide, "Every learner gets Tars.", 0.78, 1.6, 7.1, 0.62, size=32, color="5EEAD4", bold=True)
    add_text(
        slide,
        "A tutor that lives beside the cursor, sees what the learner sees, understands where they are stuck, and stays throughout the learning journey.",
        0.78,
        2.45,
        6.45,
        1.25,
        size=18,
        color=C.LIGHT_TEXT,
    )
    reveal_items = [
        ("SEES", "the learner’s current screen"),
        ("SHOWS", "exactly where to look or click"),
        ("ACTS", "inside the browser when asked"),
    ]
    for i, (label, body) in enumerate(reveal_items):
        y = 4.12 + i * 0.64
        add_text(slide, label, 0.78, y, 0.82, 0.22, size=9.5, color="5EEAD4", bold=True)
        add_text(slide, body, 1.75, y - 0.03, 4.8, 0.3, size=13.5, color=C.LIGHT_TEXT)
    add_browser_placeholder(
        slide,
        7.72,
        1.63,
        4.8,
        3.62,
        "Tars living beside the cursor",
        "Use the cleanest screenshot of the floating tutor.",
        dark=True,
    )
    add_rule(slide, 7.72, 5.78, 4.8, C.DARK_LINE)
    add_text(slide, "“It’s the attention of a personal trainer, available to every learner.”", 7.72, 6.02, 4.8, 0.66, size=14, color=C.WHITE, bold=True, italic=True)
    add_footer(slide, 5, dark=True)

    # 6 — AWS story
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.WHITE)
    add_header(slide, "05  /  The AWS moment", "The AWS lesson made sense. The UI had changed.", None, 6)
    add_text(
        slide,
        "During Cprime training, we were asked to create an EC2 instance and deploy our app. The trainer showed it once — then we had to do it ourselves.",
        0.78,
        1.62,
        11.65,
        0.68,
        size=16,
        color=C.SLATE,
    )
    timeline = [
        ("01", "We forgot where to click", "The workflow was clear in class. The exact controls were not."),
        ("02", "We jumped between ChatGPT and AWS", "The answer said “top left.” AWS had already changed the UI."),
        ("03", "With Tars: hold Ctrl, point, ask", "Tars sees this screen and points to the exact next step."),
    ]
    for i, (num, title, body) in enumerate(timeline):
        y = 2.72 + i * 1.18
        accent = C.AMBER if i < 2 else C.BLUE
        add_text(slide, num, 0.78, y, 0.42, 0.2, size=9.5, color=accent, bold=True)
        add_text(slide, title, 1.38, y - 0.03, 5.35, 0.3, size=16, color=C.INK, bold=True)
        add_text(slide, body, 1.38, y + 0.38, 5.42, 0.5, size=12.5, color=C.SLATE)
        if i < 2:
            add_rule(slide, 1.38, y + 0.94, 5.25, C.LINE)
    add_browser_placeholder(
        slide,
        7.45,
        2.68,
        5.08,
        3.18,
        "Tars pointing to Launch instance",
        "Show the exact AWS control Tars grounded to.",
    )
    add_box(slide, 0.78, 6.2, 11.75, 0.52, fill=C.BLUE_SOFT)
    add_text(slide, "No screenshot loop. No stale directions. No guessing.", 1.05, 6.36, 11.18, 0.22, size=13.5, color=C.BLUE_DARK, bold=True)
    add_footer(slide, 6)

    # 7 — Ctrl interaction
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.BG)
    add_header(slide, "06  /  The interaction", "Hold Ctrl. Point. Ask.", "The learner refers to the screen naturally. Tars responds in context.", 7)
    capabilities = [
        ("Point, circle, or underline", "Tars understands the exact region the learner means."),
        ("Ask “where do I go next?”", "Tars points to the control on the current screen."),
        ("Need a visual explanation", "Tars draws directly over the page."),
        ("Get stuck mid-workflow", "Tars can scroll, switch tabs, and navigate with permission."),
    ]
    for i, (title, body) in enumerate(capabilities):
        y = 2.08 + i * 0.98
        add_text(slide, f"0{i + 1}", 0.78, y, 0.42, 0.2, size=9.5, color=C.BLUE, bold=True)
        add_text(slide, title, 1.38, y - 0.04, 4.9, 0.3, size=15.5, color=C.INK, bold=True)
        add_text(slide, body, 1.38, y + 0.35, 4.95, 0.42, size=12.2, color=C.SLATE)
    add_browser_placeholder(
        slide,
        7.1,
        2.03,
        5.42,
        3.36,
        "Ctrl gesture, circle, or screen drawing",
        "Show the trail and the part of the screen Tars understands.",
    )
    add_text(slide, "WORKS INSIDE", 7.1, 5.78, 1.1, 0.2, size=9, color=C.MUTED, bold=True)
    platforms = ["AWS", "JIRA", "GITHUB", "SERVICENOW", "POWER BI"]
    x = 7.1
    for platform in platforms:
        width = max(0.72, len(platform) * 0.095 + 0.34)
        add_box(slide, x, 6.12, width, 0.34, fill=C.WHITE, line=C.LINE, radius=True)
        add_text(slide, platform, x, 6.22, width, 0.13, size=8.2, color=C.SLATE, bold=True, align=PP_ALIGN.CENTER)
        x += width + 0.16
    add_footer(slide, 7)

    # 8 — Live Classroom
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.WHITE)
    add_header(slide, "07  /  Live Classroom Mode", "A live classroom for one learner.", "The feel of a real tutor, with attention focused on the person in front of it.", 8)
    classroom = [
        ("REGULAR MODE", "Read lessons, explore examples, take quizzes, and ask quick questions."),
        ("LIVE CLASSROOM", "Talk naturally while Tars explains on a whiteboard and draws diagrams in real time."),
        ("IMPLEMENTATION", "Tars watches the learner apply the lesson, steps in when needed, and never judges the mistake."),
    ]
    for i, (label, body) in enumerate(classroom):
        y = 2.02 + i * 1.2
        add_text(slide, label, 0.78, y, 1.95, 0.2, size=9.5, color=C.TEAL if i else C.BLUE, bold=True)
        add_text(slide, body, 0.78, y + 0.36, 6.1, 0.78, size=14.2, color=C.INK, bold=i == 2)
        if i < 2:
            add_rule(slide, 0.78, y + 1.02, 6.12, C.LINE)
    add_browser_placeholder(
        slide,
        7.42,
        2.02,
        5.1,
        3.55,
        "Live Classroom whiteboard",
        "Show Tars teaching, drawing, or explaining live.",
    )
    add_box(slide, 0.78, 5.92, 11.74, 0.7, fill=C.TEAL_SOFT)
    add_text(slide, "Ask twice. Make mistakes. Learn without judgment.", 1.04, 6.13, 11.2, 0.26, size=15.5, color=C.TEAL, bold=True)
    add_footer(slide, 8)

    # 9 — Practice and proof
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.BG)
    add_header(slide, "08  /  From learning to proof", "It stays when the lesson ends.", "Tars stays while the learner implements what they just learned inside the real platform.", 9)
    steps = [
        ("01", "LEARN", "Source-backed lesson and explanation."),
        ("02", "PRACTICE", "A browser lab opens the real platform."),
        ("03", "OBSERVE", "Tars monitors progress and collects evidence."),
        ("04", "VERIFY", "The backend confirms the task and cleanup."),
    ]
    for i, (num, label, body) in enumerate(steps):
        x = 0.78 + i * 3.0
        add_text(slide, num, x, 2.2, 0.42, 0.2, size=9.5, color=C.BLUE, bold=True)
        add_text(slide, label, x, 2.63, 2.48, 0.28, size=15, color=C.INK, bold=True)
        add_text(slide, body, x, 3.16, 2.48, 0.72, size=12.5, color=C.SLATE)
        if i < 3:
            add_text(slide, "→", x + 2.48, 2.68, 0.38, 0.25, size=15, color=C.LINE, bold=True, align=PP_ALIGN.CENTER)
    add_rule(slide, 0.78, 4.34, 11.74)
    add_text(slide, "A lab completes only when the task is correct", 0.78, 4.82, 5.95, 0.38, size=17, color=C.INK, bold=True)
    add_text(slide, "and every required cleanup step is confirmed.", 0.78, 5.32, 5.9, 0.34, size=16, color=C.BLUE, bold=True)
    add_text(slide, "Learners move from", 7.48, 4.82, 2.2, 0.24, size=10, color=C.MUTED, bold=True)
    add_text(slide, "“I watched the training”", 7.48, 5.2, 4.7, 0.34, size=17, color=C.SLATE, italic=True)
    add_text(slide, "to  “I can actually do the work.”", 7.48, 5.72, 4.72, 0.38, size=17, color=C.TEAL, bold=True)
    add_footer(slide, 9)

    # 10 — Impact
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.WHITE)
    add_header(slide, "09  /  Impact & benefit", "Cprime’s expertise becomes a product it owns.", "Expert-led learning remains the foundation. Ctrl+Teach adds a scalable layer on top.", 10)
    add_text(slide, "Here’s the uncomfortable truth:", 0.78, 2.03, 5.6, 0.32, size=16, color=C.AMBER, bold=True)
    add_text(
        slide,
        "A training business that grows only by adding trainers will always be limited by people, schedules, and delivery costs.",
        0.78,
        2.55,
        5.72,
        1.25,
        size=20,
        color=C.INK,
        bold=True,
    )
    add_text(
        slide,
        "Tools change faster than traditional courses can be updated, and learners cannot wait for the next classroom session when they get stuck.",
        0.78,
        4.12,
        5.72,
        0.92,
        size=14.5,
        color=C.SLATE,
    )
    outcomes = [
        ("LOWER DELIVERY EFFORT", "Support more learners without delivery capacity growing at the same rate."),
        ("RECURRING PRODUCT REVENUE", "Sell a product Cprime owns instead of rebuilding every experience from scratch."),
        ("FASTER COURSE UPDATES", "Improve the core once and carry the change into future learning journeys."),
        ("INDIVIDUAL ATTENTION AT SCALE", "Give each learner memory, context, and help at the moment of need."),
    ]
    for i, (label, body) in enumerate(outcomes):
        y = 2.0 + i * 1.06
        add_text(slide, f"0{i + 1}", 7.04, y, 0.42, 0.2, size=9.5, color=C.BLUE, bold=True)
        add_text(slide, label, 7.64, y - 0.02, 4.6, 0.22, size=9.8, color=C.BLUE, bold=True)
        add_text(slide, body, 7.64, y + 0.34, 4.55, 0.55, size=12.5, color=C.SLATE)
    add_box(slide, 0.78, 6.15, 11.74, 0.58, fill=C.NAVY)
    add_text(slide, "Created once. Personalized automatically. Sold repeatedly.", 1.05, 6.32, 11.2, 0.22, size=15.5, color=C.WHITE, bold=True)
    add_footer(slide, 10)

    # 11 — Product leverage
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.BG)
    add_header(slide, "10  /  Product leverage", "Build it once. Sell it thousands of times.", "A product that carries Cprime’s expertise further — and compounds with every improvement.", 11)
    leverage = [
        ("01", "CREATE", "Cprime’s experts shape the knowledge, standards, and learning outcomes."),
        ("02", "PERSONALIZE", "Ctrl+Teach adapts the course and tutor to the learner’s memory and progress."),
        ("03", "COMPOUND", "Every content update and tutor improvement benefits the next learner."),
    ]
    for i, (num, label, body) in enumerate(leverage):
        x = 0.78 + i * 4.0
        if i:
            add_box(slide, x - 0.28, 2.12, 0.012, 2.34, fill=C.LINE)
        add_text(slide, num, x, 2.12, 0.42, 0.2, size=9.5, color=C.BLUE, bold=True)
        add_text(slide, label, x, 2.68, 3.4, 0.32, size=17, color=C.INK, bold=True)
        add_text(slide, body, x, 3.3, 3.46, 1.02, size=14, color=C.SLATE)
    add_rule(slide, 0.78, 4.86, 11.74)
    add_text(slide, "This could become Cprime’s first truly scalable learning product.", 0.78, 5.26, 10.9, 0.42, size=22, color=C.BLUE, bold=True)
    add_text(slide, "Not another video library — a tutor carrying Cprime’s knowledge and staying until the learner can do it alone.", 0.78, 5.95, 11.3, 0.58, size=15, color=C.SLATE)
    add_footer(slide, 11)

    # 12 — Close
    slide = prs.slides.add_slide(blank)
    set_background(slide, C.NAVY)
    add_box(slide, 0.78, 0.62, 0.06, 0.3, fill=C.TEAL)
    add_text(slide, "CTRL+TEACH", 0.98, 0.62, 3.2, 0.28, size=10.5, color="5EEAD4", bold=True)
    add_text(slide, "The future of learning will not be", 0.78, 1.62, 11.4, 0.58, size=32, color=C.WHITE, bold=True)
    add_text(slide, "another video library.", 0.78, 2.28, 11.4, 0.58, size=32, color="60A5FA", bold=True)
    add_rule(slide, 0.78, 3.22, 11.72, C.DARK_LINE)
    add_text(
        slide,
        "It will be a tutor that stays beside you until you can do it yourself.",
        0.78,
        3.75,
        10.8,
        1.05,
        size=28,
        color=C.WHITE,
        bold=True,
    )
    add_text(slide, "Courses that adapt.", 0.78, 5.45, 2.8, 0.28, size=13, color="B9C0CC", bold=True)
    add_text(slide, "A tutor that acts.", 3.72, 5.45, 2.8, 0.28, size=13, color="B9C0CC", bold=True)
    add_text(slide, "Learning that proves itself.", 6.62, 5.45, 3.5, 0.28, size=13, color="5EEAD4", bold=True)
    add_text(slide, "Thank you.", 0.78, 6.64, 2.0, 0.28, size=11.5, color="98A2B3", bold=True)

    prs.core_properties.title = "Ctrl+Teach — Clean Modern Hackathon Pitch"
    prs.core_properties.subject = "Cprime scalable learning product pitch"
    prs.core_properties.author = "Ctrl+Teach"
    prs.save(OUT)
    print(f"Wrote {OUT.resolve()} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    build_deck()
